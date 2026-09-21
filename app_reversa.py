from datetime import datetime, timedelta, timezone
import io
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
import pandas as pd
import streamlit as st

# Configuração da Página & Tema Logístico
st.set_page_config(
    page_title="WMS & Logística Reversa - RMC Mariano",
    page_icon="🏭",
    layout="wide",
)

# Arquivo local para salvar os dados
DB_FILE = "dados_reversas.csv"


# Função para pegar o horário exato de Brasília (-3 horas do UTC) sem erros de biblioteca
def obter_horario_brasilia():
  fuso_brasilia = timezone(timedelta(hours=-3))
  return datetime.now(fuso_brasilia).strftime("%d/%m/%Y %H:%M")


# Função para carregar os dados com segurança e tipos corretos
def carregar_dados():
  colunas_necessarias = [
      "ID_Devolucao",
      "Data_Registro",
      "Pedido",
      "NF",
      "Cliente",
      "Cidade",
      "Transportadora",
      "Motivo",
      "Itens",
      "Status",
      "Data_Inicio_Coleta",
      "Entregador_Responsavel",
      "Observacao_Transportadora",  # Nova coluna adicionada
      "Evidencia_Anexo",
      "Data_Conclusao",
  ]

  if os.path.exists(DB_FILE):
    df = pd.read_csv(DB_FILE, dtype=str)
    for col in colunas_necessarias:
      if col not in df.columns:
        df[col] = ""
      else:
        df[col] = df[col].fillna("")
    return df
  else:
    df_padrao = pd.DataFrame(columns=colunas_necessarias, dtype=str)
    return df_padrao


# Função para salvar os dados
def salvar_dados(df):
  df.to_csv(DB_FILE, index=False)


# Carrega o banco de dados na sessão
if "df_reversas" not in st.session_state:
  st.session_state.df_reversas = carregar_dados()

# Cabeçalho com identidade visual de Logística / Centro de Distribuição
st.markdown(
    """
    <div style="background-color: #0f172a; padding: 20px; border-radius: 8px; border-left: 6px solid #f59e0b;">
        <h2 style="color: white; margin: 0; font-family: sans-serif;">🏭 RMC MARIANO | Gestão de Logística Reversa & WMS</h2>
        <p style="color: #94a3b8; margin: 5px 0 0 0; font-size: 14px;">Controle de fluxo de coletas, rastreio de transportadora e monitoramento de status.</p>
    </div>
""",
    unsafe_allow_html=True,
)
st.markdown("<br>", unsafe_allow_html=True)

# Criação das Abas Principais
aba_gestor, aba_transportadora, aba_analise, aba_relatorios, aba_termo = st.tabs([
    "📋 1. Painel Operacional (Cadastro & Histórico)",
    "🚚 2. Portal da Transportadora",
    "🔍 3. Analisar Pedidos Finalizados",
    "📊 4. Dashboard & Indicadores",
    "📄 5. Emissão de Termos (Word)",
])

# ==========================================
# ABA 1: PAINEL OPERACIONAL (CADASTRO E HISTÓRICO DE STATUS)
# ==========================================
with aba_gestor:
  st.subheader("📝 Abertura de Ordem de Coleta Reversa")
  st.markdown(
      "Cadastre a ordem de reversa. O pedido será enviado diretamente para a"
      " transportadora realizar a verificação e coleta."
  )

  with st.form("form_cadastro_reversa", clear_on_submit=True):
    col1, col2, col3 = st.columns(3)

    with col1:
      num_pedido = st.text_input("Nº do Pedido / Carga")
      num_nf = st.text_input("Nota Fiscal (NF)")

    with col2:
      cliente = st.text_input("Nome da Revendedora / Cliente Final")
      cidade = st.text_input("Destino / Cidade")

    with col3:
      transportadora = "Transportadora José Augusto"
      st.text_input(
          "Operador Logístico / Transportadora",
          value=transportadora,
          disabled=True,
      )
      motivo = st.selectbox(
          "Motivo da Ocorrência",
          ["Produto Avariado", "Troca de Caixa", "Desistência do Pedido"],
      )

    itens = st.text_area(
        "SKUs / Descrição dos Itens a Serem Coletados",
        placeholder="Ex: 1x Colônia Floratta Red, 2x Hidratante Instance...",
    )

    submitted = st.form_submit_button(
        "🚀 Emitir Ordem e Enviar para Transportadora",
        use_container_width=True,
    )

    if submitted:
      if not num_pedido or not cliente:
        st.warning(
            "⚠️ Preencha os campos obrigatórios (Nº do Pedido e Cliente)!"
        )
      else:
        fuso_brasilia = timezone(timedelta(hours=-3))
        novo_id = f"REV-{datetime.now(fuso_brasilia).strftime('%Y%m%d%H%M%S')}"
        nova_linha = {
            "ID_Devolucao": str(novo_id),
            "Data_Registro": obter_horario_brasilia(),
            "Pedido": str(num_pedido),
            "NF": str(num_nf) if num_nf else "N/A",
            "Cliente": str(cliente).upper(),
            "Cidade": str(cidade).upper(),
            "Transportadora": str(transportadora),
            "Motivo": str(motivo),
            "Itens": str(itens),
            "Status": "🟡 Aguardando Verificação",
            "Data_Inicio_Coleta": "",
            "Entregador_Responsavel": "",
            "Observacao_Transportadora": "",
            "Evidencia_Anexo": "",
            "Data_Conclusao": "",
        }

        st.session_state.df_reversas = pd.concat(
            [st.session_state.df_reversas, pd.DataFrame([nova_linha])],
            ignore_index=True,
        )
        salvar_dados(st.session_state.df_reversas)
        st.success(
            f"✅ Ordem **{novo_id}** gerada com sucesso e enviada para o portal"
            " da transportadora!"
        )

  st.markdown("---")
  st.subheader(
      "📊 Histórico Geral de Status de Cada Pedido (Acompanhamento)"
  )
  st.markdown(
      "Visualize abaixo o andamento de todas as ordens abertas e concluídas."
  )

  df_atual = st.session_state.df_reversas
  if not df_atual.empty:
    filtro_status = st.selectbox(
        "Filtrar por Status Operacional",
        [
            "Todos",
            "🟡 Aguardando Verificação",
            "🚚 Coleta Iniciada",
            "✅ Finalizado pela Transportadora",
        ],
    )
    if filtro_status != "Todos":
      df_filtrado = df_atual[df_atual["Status"] == filtro_status]
    else:
      df_filtrado = df_atual

    st.dataframe(df_filtrado, use_container_width=True)

    st.markdown("### 🗑️ Gestão de Registros (Exclusão Direta)")
    if not df_filtrado.empty:
      for index, row in df_filtrado.iterrows():
        col_reg1, col_reg2, col_reg3 = st.columns([1.5, 4, 1])
        with col_reg1:
          st.text(row["ID_Devolucao"])
        with col_reg2:
          st.text(
              f"Ped: {row['Pedido']} | Cliente: {row['Cliente']} ("
              f"{row['Status']})"
          )
        with col_reg3:
          if st.button("🗑️ Excluir", key=f"lixeira_{row['ID_Devolucao']}"):
            st.session_state.df_reversas = st.session_state.df_reversas[
                st.session_state.df_reversas["ID_Devolucao"]
                != row["ID_Devolucao"]
            ].reset_index(drop=True)
            salvar_dados(st.session_state.df_reversas)
            st.success(f"Registro {row['ID_Devolucao']} excluído com sucesso!")
            st.rerun()
    else:
      st.info("Nenhum registro para gerenciar nesta visualização.")
  else:
    st.info("Nenhuma ordem de devolução registrada no sistema.")

# ==========================================
# ABA 2: PORTAL DA TRANSPORTADORA
# ==========================================
with aba_transportadora:
  st.subheader(
      "🚚 Portal do Operador Logístico (Verificação, Coleta e Envio de Caixa)"
  )
  st.markdown(
      "Gerencie os pedidos aguardando verificação, autorize e registre a"
      " data/hora da coleta, previsões e ocorrências."
  )

  transp_selecionada = "Transportadora José Augusto"
  df_transp = st.session_state.df_reversas[
      st.session_state.df_reversas["Transportadora"] == transp_selecionada
  ]

  cidades_disponiveis = ["Todas"] + list(df_transp["Cidade"].unique())
  filtro_cidade_transp = st.selectbox(
      "Filtrar por Rota / Cidade:", cidades_disponiveis, key="filtro_cid_transp"
  )

  if filtro_cidade_transp != "Todas":
    pendentes_transp = df_transp[
        (df_transp["Status"] != "✅ Finalizado pela Transportadora")
        & (df_transp["Cidade"] == filtro_cidade_transp)
    ]
  else:
    pendentes_transp = df_transp[
        df_transp["Status"] != "✅ Finalizado pela Transportadora"
    ]

  st.markdown("---")
  st.markdown(
      f"### 📦 Ordens Ativas no Painel ({len(pendentes_transp)} ordens)"
  )

  if not pendentes_transp.empty:
    for index, row in pendentes_transp.iterrows():
      with st.container(border=True):
        col_a, col_b = st.columns([2.5, 1.5])
        with col_a:
          st.markdown(
              f"**ID:** `{row['ID_Devolucao']}` | **Status Atual:**"
              f" **{row['Status']}**"
          )
          st.markdown(
              f"📄 **Pedido:** `{row['Pedido']}` | **NF:** `{row['NF']}`"
          )
          st.markdown(
              f"👤 **Cliente:** **{row['Cliente']}** (📍 {row['Cidade']})"
          )
          st.markdown(f"❓ **Motivo:** `{row['Motivo']}`")
          st.markdown(f"📦 **Itens a Coletar:** {row['Itens']}")
          if row["Data_Inicio_Coleta"]:
            st.markdown(
                f"⏱️ **Autorizado/Início da Coleta:**"
                f" `{row['Data_Inicio_Coleta']}` por"
                f" `{row['Entregador_Responsavel']}`"
            )
          if row["Observacao_Transportadora"]:
            st.markdown(
                f"💬 **Obs/Previsão da Transportadora:**"
                f" `{row['Observacao_Transportadora']}`"
            )

        with col_b:
          st.markdown("##### ✍️ Ações da Transportadora")

          # Ação 1: Iniciar Coleta
          if row["Status"] == "🟡 Aguardando Verificação":
            nome_entregador = st.text_input(
                "Motorista Responsável:",
                key=f"ent_{row['ID_Devolucao']}",
                placeholder="Nome do motorista",
            )
            obs_transp_inicial = st.text_area(
                "Previsão / Observações da Rota:",
                key=f"obs_ini_{row['ID_Devolucao']}",
                placeholder="Ex: Previsão de coleta para amanhã às 14h...",
            )

            if st.button(
                "🚀 Autorizar e Iniciar Coleta",
                key=f"btn_iniciar_{row['ID_Devolucao']}",
            ):
              if not nome_entregador.strip():
                st.error("⚠️ Informe o nome do motorista!")
              else:
                idx_real = st.session_state.df_reversas[
                    st.session_state.df_reversas["ID_Devolucao"]
                    == row["ID_Devolucao"]
                ].index[0]
                st.session_state.df_reversas.at[
                    idx_real, "Status"
                ] = "🚚 Coleta Iniciada"
                st.session_state.df_reversas.at[
                    idx_real, "Data_Inicio_Coleta"
                ] = obter_horario_brasilia()
                st.session_state.df_reversas.at[
                    idx_real, "Entregador_Responsavel"
                ] = str(nome_entregador).upper()
                st.session_state.df_reversas.at[
                    idx_real, "Observacao_Transportadora"
                ] = str(obs_transp_inicial)
                salvar_dados(st.session_state.df_reversas)
                st.success(
                    "Coleta autorizada e horário oficial registrado com"
                    " sucesso!"
                )
                st.rerun()

          # Ação 2: Finalizar Pedido com anexo de foto/vídeo e atualização de obs
          if row["Status"] == "🚚 Coleta Iniciada":
            obs_transp_final = st.text_area(
                "Atualizar Previsão / Relatório de Conclusão:",
                value=row["Observacao_Transportadora"],
                key=f"obs_fin_{row['ID_Devolucao']}",
                placeholder=(
                    "Ex: Coletado com sucesso, caixa em trânsito para o CD..."
                ),
            )

            uploaded_file = st.file_uploader(
                "Anexar Foto ou Vídeo (Evidência):",
                type=["png", "jpg", "jpeg", "mp4", "mov"],
                key=f"file_{row['ID_Devolucao']}",
            )

            if st.button(
                "✅ Finalizar Pedido (Enviar Caixa)",
                key=f"btn_finalizar_{row['ID_Devolucao']}",
            ):
              idx_real = st.session_state.df_reversas[
                  st.session_state.df_reversas["ID_Devolucao"]
                  == row["ID_Devolucao"]
              ].index[0]

              nome_arquivo = ""
              if uploaded_file is not None:
                nome_arquivo = uploaded_file.name
                os.makedirs("uploads", exist_ok=True)
                with open(
                    os.path.join("uploads", uploaded_file.name), "wb"
                ) as f:
                  f.write(uploaded_file.getbuffer())

              st.session_state.df_reversas.at[
                  idx_real, "Status"
              ] = "✅ Finalizado pela Transportadora"
              st.session_state.df_reversas.at[
                  idx_real, "Observacao_Transportadora"
              ] = str(obs_transp_final)
              st.session_state.df_reversas.at[
                  idx_real, "Evidencia_Anexo"
              ] = nome_arquivo
              st.session_state.df_reversas.at[
                  idx_real, "Data_Conclusao"
              ] = obter_horario_brasilia()

              salvar_dados(st.session_state.df_reversas)
              st.success(
                  "Pedido finalizado e enviado para análise no Centro de"
                  " Distribuição!"
              )
              st.rerun()
  else:
    st.info(
        "🎉 Nenhuma ordem ativa pendente na transportadora para as rotas"
        " selecionadas."
    )

# ==========================================
# ABA 3: ANALISAR PEDIDOS FINALIZADOS (COM VISUALIZAÇÃO DE ANEXOS)
# ==========================================
with aba_analise:
  st.subheader("🔍 Central de Análise de Pedidos Finalizados")
  st.markdown(
      "Aqui você visualiza os dados, confere o histórico, as previsões da"
      " transportadora e as evidências (fotos/vídeos)."
  )

  df_geral = st.session_state.df_reversas
  df_finalizados = df_geral[
      df_geral["Status"] == "✅ Finalizado pela Transportadora"
  ]

  if not df_finalizados.empty:
    st.markdown(
        f"### 📥 Total de pedidos aguardando sua análise:"
        f" {len(df_finalizados)}"
    )

    for index, row in df_finalizados.iterrows():
      with st.container(border=True):
        col_an1, col_an2 = st.columns([2, 1.5])
        with col_an1:
          st.markdown(
              f"**ID:** `{row['ID_Devolucao']}` | **Conclusão:**"
              f" `{row['Data_Conclusao']}`"
          )
          st.markdown(
              f"📄 **Pedido:** `{row['Pedido']}` | **NF:** `{row['NF']}`"
          )
          st.markdown(
              f"👤 **Cliente:** **{row['Cliente']}** (📍 {row['Cidade']})"
          )
          st.markdown(f"❓ **Motivo:** `{row['Motivo']}`")
          st.markdown(f"📦 **Itens:** {row['Itens']}")
          st.markdown(
              f"⏱️ **Início/Autorização da Coleta:**"
              f" `{row['Data_Inicio_Coleta']}`"
          )
          st.markdown(
              f"🚚 **Motorista Responsável:** `{row['Entregador_Responsavel']}`"
          )
          if row["Observacao_Transportadora"]:
            st.markdown(
                f"💬 **Relatório/Previsão da Transportadora:**"
                f" `{row['Observacao_Transportadora']}`"
            )

        with col_an2:
          st.markdown("##### 📎 Evidência Anexada")
          arquivo_anexado = row["Evidencia_Anexo"]

          if arquivo_anexado:
            st.success(f"Arquivo: {arquivo_anexado}")
            caminho_arquivo = os.path.join("uploads", arquivo_anexado)

            if os.path.exists(caminho_arquivo):
              extensao = arquivo_anexado.split(".")[-1].lower()
              if extensao in ["png", "jpg", "jpeg"]:
                st.image(
                    caminho_arquivo,
                    caption="Foto enviada pelo entregador",
                    use_container_width=True,
                )
              elif extensao in ["mp4", "mov"]:
                st.video(caminho_arquivo)

              with open(caminho_arquivo, "rb") as file_to_download:
                st.download_button(
                    label="📥 Baixar Evidência",
                    data=file_to_download,
                    file_name=arquivo_anexado,
                    key=f"dl_{row['ID_Devolucao']}",
                )
            else:
              st.warning(
                  "⚠️ O registro indica um anexo, mas o arquivo físico não foi"
                  " encontrado na pasta do sistema."
              )
          else:
            st.info(
                "Nenhum arquivo de foto ou vídeo foi anexado pela transportadora"
                " neste pedido."
            )
  else:
    st.info(
        "Nenhum pedido finalizado pela transportadora aguardando análise no"
        " momento."
    )

# ==========================================
# ABA 4: INDICADORES & ALERTAS
# ==========================================
with aba_relatorios:
  st.subheader("📊 Dashboard de Performance Logística (KPIs)")
  df_geral = st.session_state.df_reversas

  if not df_geral.empty:
    total_geral = len(df_geral)
    total_aguardando = len(
        df_geral[df_geral["Status"] == "🟡 Aguardando Verificação"]
    )
    total_coleta_iniciada = len(
        df_geral[df_geral["Status"] == "🚚 Coleta Iniciada"]
    )
    total_finalizados = len(
        df_geral[df_geral["Status"] == "✅ Finalizado pela Transportadora"]
    )

    col_m1, col_m2, col_m3, col_m4 = st.columns(4)
    col_m1.metric("Total de Ordens", total_geral)
    col_m2.metric("Aguardando Verificação", total_aguardando)
    col_m3.metric("Coleta Iniciada", total_coleta_iniciada)
    col_m4.metric("Finalizados (Para Análise)", total_finalizados)
  else:
    st.info("Insira dados para visualizar os indicadores do painel.")

# ==========================================
# ABA 5: GERADOR DE TERMO / CHECKLIST (PADRÃO ANTIGO + ITENS EM CHECKLIST)
# ==========================================
with aba_termo:
  st.subheader("📄 Geração de Documentação Oficial (Termo & Checklist)")
  st.markdown(
      "Emita o termo padrão de responsabilidade com a lista de itens convertida"
      " em checklist de conferência."
  )

  df_termo_geral = st.session_state.df_reversas

  if not df_termo_geral.empty:
    lista_ids = df_termo_geral["ID_Devolucao"].tolist()
    id_escolhido = st.selectbox(
        "Selecione o ID da Ordem para Gerar o Documento:",
        lista_ids,
        key="sel_id_termo",
    )

    dados_linha = df_termo_geral[
        df_termo_geral["ID_Devolucao"] == id_escolhido
    ].iloc[0]

    st.markdown("---")
    st.markdown("#### 🔍 Prévia dos Dados para Impressão:")
    c_t1, c_t2, c_t3 = st.columns(3)
    c_t1.text_input("ID", value=dados_linha["ID_Devolucao"], disabled=True)
    c_t2.text_input("Pedido", value=dados_linha["Pedido"], disabled=True)
    c_t3.text_input("Nota Fiscal", value=dados_linha["NF"], disabled=True)

    st.text_input("Cliente", value=dados_linha["Cliente"], disabled=True)
    st.text_input("Motivo", value=dados_linha["Motivo"], disabled=True)
    st.text_area("Itens", value=dados_linha["Itens"], disabled=True)
    if dados_linha["Observacao_Transportadora"]:
      st.text_area(
          "Obs da Transportadora",
          value=dados_linha["Observacao_Transportadora"],
          disabled=True,
      )

    if st.button(
        "📥 Baixar Termo em Formato Word (.docx)",
        use_container_width=True,
        key="btn_dl_termo",
    ):
      doc = Document()

      if os.path.exists("logo.png"):
        try:
          doc.add_picture("logo.png", width=Inches(2.0))
        except:
          pass

      # Título do Termo Original
      p_titulo = doc.add_paragraph()
      run_titulo = p_titulo.add_run(
          "GRUPO RMC MARIANO - TERMO DE RESPONSABILIDADE E CHECKLIST DE"
          " DEVOLUÇÃO"
      )
      run_titulo.bold = True
      run_titulo.font.size = Pt(13)

      doc.add_paragraph(f"Data de Emissão: {dados_linha['Data_Registro']}")
      doc.add_paragraph(f"ID da Ocorrência: {dados_linha['ID_Devolucao']}")
      doc.add_paragraph(f"Nº do Pedido: {dados_linha['Pedido']}")
      doc.add_paragraph(f"Nota Fiscal: {dados_linha['NF']}")
      doc.add_paragraph(f"Revendedora / Cliente: {dados_linha['Cliente']}")
      doc.add_paragraph(f"Cidade: {dados_linha['Cidade']}")
      doc.add_paragraph(f"Motivo da Devolução: {dados_linha['Motivo']}")

      if dados_linha["Observacao_Transportadora"]:
        doc.add_paragraph(
            f"Observação / Previsão da Transportadora:"
            f" {dados_linha['Observacao_Transportadora']}"
        )

      doc.add_paragraph(
          "\nITENS A SEREM CONFERIDOS E COLETADOS (CHECKLIST DE RECEBIMENTO):"
      )

      # Transforma os itens digitados pelo usuário em linhas de checklist com caixas [ ]
      texto_itens = str(dados_linha["Itens"]).strip()
      if texto_itens:
        linhas_itens = [
            linha.strip()
            for linha in texto_itens.replace(",", "\n").split("\n")
            if linha.strip()
        ]

        tabela_itens = doc.add_table(rows=len(linhas_itens) + 1, cols=2)
        tabela_itens.alignment = WD_TABLE_ALIGNMENT.CENTER

        hdr_cells = tabela_itens.rows[0].cells
        hdr_cells[0].text = "Descrição do Item / SKU"
        hdr_cells[1].text = "STATUS ([ ] OK  /  [ ] FALTA)"
        hdr_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        for cell in hdr_cells:
          shd = parse_xml(r'<w:shd {} w:fill="1A5276"/>'.format(nsdecls("w")))
          cell._tc.get_or_add_tcPr().append(shd)
          for paragraph in cell.paragraphs:
            for run in paragraph.runs:
              run.font.color.rgb = RGBColor(255, 255, 255)
              run.bold = True

        for idx, item_txt in enumerate(linhas_itens):
          r_cells = tabela_itens.rows[idx + 1].cells
          r_cells[0].text = item_txt
          r_cells[1].text = "[   ] CONFERIDO      [   ] AVARIADO"
          r_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
      else:
        doc.add_paragraph("[Nenhum item específico detalhado]")

      doc.add_paragraph(
          "\nDeclaro para os devidos fins que os produtos acima descritos"
          " estão sendo entregues à Transportadora José Augusto para retorno"
          " ao Centro de Distribuição do Grupo RMC Mariano."
      )

      doc.add_paragraph("\n\n__________________________________________________")
      doc.add_paragraph(f"Assinatura da Revendedora: {dados_linha['Cliente']}")

      doc.add_paragraph("__________________________________________________")
      doc.add_paragraph("Assinatura / Carimbo da Transportadora José Augusto")

      buffer = io.BytesIO()
      doc.save(buffer)
      buffer.seek(0)

      st.download_button(
          label="💾 Clique aqui para baixar o documento Word gerado",
          data=buffer,
          file_name=f"Termo_Devolucao_{dados_linha['ID_Devolucao']}.docx",
          mime=(
              "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
          ),
      )
      st.success("Documento gerado com sucesso!")
  else:
    st.info("Insira ordens nas abas anteriores para gerar documentos.")
